# =============================================================================
# core/indexing.py
# WHAT THIS FILE DOES: Keeps the Chroma index synchronized with documents/
#
# FLOW:
#   documents/ folder
#       --> hash supported files to detect new/changed content
#       --> load and split files that need indexing
#       --> remove chunks for deleted or changed files
#       --> add the current chunks to Chroma
#       --> save a tracker matching the current folder
# =============================================================================

# Step 1: Import file, hashing, document-loading, splitting, and Chroma helpers.
import hashlib
import json
import os
from datetime import date, datetime

from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_community.vectorstores import Chroma
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from core.embeddings import get_embeddings

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TRACKER_PATH = os.path.join(BASE_DIR, "index_tracker.json")
SUPPORTED_EXTENSIONS = (".txt", ".md", ".pdf", ".xls", ".xlsx")


# Step 2: Normalize paths and load/hash files using consistent source metadata.
def _normalized_path(filepath: str) -> str:
    """Make document paths comparable across Chroma metadata and the filesystem."""
    return os.path.normcase(os.path.abspath(filepath))


def _file_digest(filepath: str) -> str:
    """Hash file contents so same-size edits and timestamp changes are detected."""
    digest = hashlib.sha256()
    with open(filepath, "rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _format_cell_value(value: object) -> str:
    """Convert spreadsheet cell values into readable text for document search."""
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def _spreadsheet_document(
    filepath: str,
    sheet_name: str,
    rows: list[tuple],
) -> Document:
    """Use the first populated row as headers and preserve workbook/sheet metadata."""
    source = _normalized_path(filepath)
    populated_rows = [
        (row_number, [_format_cell_value(value) for value in row])
        for row_number, row in enumerate(rows, start=1)
        if any(_format_cell_value(value) for value in row)
    ]

    if not populated_rows:
        content = ""
    else:
        _, headers = populated_rows[0]
        column_count = max(len(values) for _, values in populated_rows)
        column_headers = [
            (
                headers[column_number - 1]
                if column_number <= len(headers) and headers[column_number - 1]
                else f"Column {column_number}"
            )
            for column_number in range(1, column_count + 1)
        ]
        lines = [
            f"Worksheet: {sheet_name}",
            "Columns: " + " | ".join(column_headers),
        ]
        for row_number, values in populated_rows[1:]:
            cells = [
                f"{column_headers[column_number - 1]}: {value}"
                for column_number, value in enumerate(values, start=1)
                if value
            ]
            if cells:
                lines.append(f"Row {row_number}: " + " | ".join(cells))
        content = "\n".join(lines)

    return Document(
        page_content=content,
        metadata={"source": source, "sheet": sheet_name},
    )


def _load_xlsx(filepath: str) -> list[Document]:
    """Read every .xlsx worksheet, using saved formula results when available."""
    from openpyxl import load_workbook

    workbook = load_workbook(filepath, read_only=True, data_only=True)
    try:
        return [
            _spreadsheet_document(
                filepath,
                worksheet.title,
                list(worksheet.iter_rows(values_only=True)),
            )
            for worksheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _load_xls(filepath: str) -> list[Document]:
    """Read every worksheet in a legacy .xls workbook."""
    import xlrd

    workbook = xlrd.open_workbook(filepath, on_demand=True)
    try:
        documents = []
        for worksheet in workbook.sheets():
            rows = []
            for row_number in range(worksheet.nrows):
                values = []
                for cell in worksheet.row(row_number):
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate.xldate_as_datetime(
                            cell.value, workbook.datemode
                        )
                    elif cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        value = None
                    else:
                        value = cell.value
                    values.append(value)
                rows.append(tuple(values))
            documents.append(
                _spreadsheet_document(filepath, worksheet.name, rows)
            )
        return documents
    finally:
        workbook.release_resources()


def _load_document(filepath: str) -> list[Document]:
    """Load a supported file and normalize its source metadata for safe cleanup."""
    source = _normalized_path(filepath)
    extension = os.path.splitext(filepath)[1].lower()

    if extension == ".pdf":
        documents = PyPDFLoader(filepath).load()
    elif extension == ".xlsx":
        return _load_xlsx(filepath)
    elif extension == ".xls":
        return _load_xls(filepath)
    else:
        documents = TextLoader(filepath, encoding="utf-8").load()

    # Preserve page metadata while normalizing the source used for cleanup.
    for document in documents:
        document.metadata["source"] = source
    return documents


def load_tracker() -> dict:
    """Read the index manifest; legacy timestamp entries are upgraded on reindex."""
    if not os.path.exists(TRACKER_PATH):
        return {}

    with open(TRACKER_PATH, "r", encoding="utf-8") as tracker_file:
        tracker = json.load(tracker_file)

    if not isinstance(tracker, dict):
        raise ValueError(f"Index tracker must contain a JSON object: {TRACKER_PATH}")
    return tracker


def save_tracker(tracker: dict) -> None:
    """Write the manifest only after Chroma has been updated successfully."""
    with open(TRACKER_PATH, "w", encoding="utf-8") as tracker_file:
        json.dump(tracker, tracker_file, indent=2)


def build_index(config: dict) -> tuple[Chroma, list[str], list[str]]:
    """
    Synchronize Chroma with the supported files currently in documents/.

    Returns the vector store, changed/new filenames, and files without readable
    text. Old timestamp-only tracker entries are safely reindexed.
    """
    # Step 3: Find supported files and connect to the configured Chroma collection.
    documents_folder = config["documents_folder"]
    os.makedirs(documents_folder, exist_ok=True)

    tracker = load_tracker()
    current_files = {}
    for filename in sorted(os.listdir(documents_folder)):
        filepath = os.path.join(documents_folder, filename)
        if (
            os.path.isfile(filepath)
            and os.path.splitext(filename)[1].lower() in SUPPORTED_EXTENSIONS
        ):
            current_files[filename] = filepath

    vectorstore = Chroma(
        persist_directory=config["chroma_db_path"],
        embedding_function=None,
        collection_name=config["collection_name"],
    )

    # Step 4: Compare the current files with both the tracker and stored chunks.
    # Inspect indexed source metadata as well as the tracker. This detects a
    # manually cleared Chroma database even when the tracker still exists.
    stored_data = vectorstore.get(include=["metadatas"])
    stored_ids_by_source = {}
    ids_without_source = []
    for chunk_id, metadata in zip(stored_data["ids"], stored_data["metadatas"]):
        source = metadata.get("source") if metadata else None
        if source:
            normalized_source = _normalized_path(source)
            stored_ids_by_source.setdefault(normalized_source, []).append(chunk_id)
        else:
            # Chunks without a source cannot be cited or synchronized reliably.
            ids_without_source.append(chunk_id)

    current_sources = {
        _normalized_path(filepath) for filepath in current_files.values()
    }
    current_digests = {
        filename: _file_digest(filepath)
        for filename, filepath in current_files.items()
    }
    files_to_index = []
    newly_indexed = []

    for filename, filepath in current_files.items():
        old_record = tracker.get(filename)
        source = _normalized_path(filepath)
        digest = current_digests[filename]

        # Old timestamp-only entries have no content hash, so reindex them once.
        changed = not isinstance(old_record, dict) or old_record.get("sha256") != digest
        expected_chunk_count = (
            old_record.get("chunk_count", 0)
            if isinstance(old_record, dict)
            else 0
        )
        missing_from_chroma = expected_chunk_count > 0 and source not in stored_ids_by_source

        if changed or missing_from_chroma:
            files_to_index.append((filename, filepath))
            newly_indexed.append(filename)

    # Step 5: Load changed documents and split their text into searchable chunks.
    sources_to_replace = {
        _normalized_path(filepath) for _, filepath in files_to_index
    }

    # Load failures are allowed to surface instead of silently keeping bad index
    # state; the UI will show the error and the tracker will not be advanced.
    documents = []
    files_without_text = []
    for filename, filepath in files_to_index:
        loaded_documents = _load_document(filepath)
        if not any(document.page_content.strip() for document in loaded_documents):
            files_without_text.append(filename)
        documents.extend(loaded_documents)

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=config["chunk_size"],
        chunk_overlap=config["chunk_overlap"],
    )
    chunks = splitter.split_documents(documents)
    # Initialize all files so empty files cannot cause a KeyError here.
    chunk_counts = {filename: 0 for filename, _ in files_to_index}
    for chunk in chunks:
        source = _normalized_path(chunk.metadata["source"])
        for filename, filepath in files_to_index:
            if _normalized_path(filepath) == source:
                chunk_counts[filename] += 1
                break

    # Step 6: Remove stale chunks, add replacements, then save the tracker.
    # Delete chunks whose source was removed, plus old chunks for changed files.
    # Deleting by stored IDs also handles legacy source-path formatting safely.
    ids_to_delete = ids_without_source + [
        chunk_id
        for source, chunk_ids in stored_ids_by_source.items()
        if source not in current_sources or source in sources_to_replace
        for chunk_id in chunk_ids
    ]
    if ids_to_delete:
        vectorstore.delete(ids=ids_to_delete)

    if chunks:
        vectorstore = Chroma(
            persist_directory=config["chroma_db_path"],
            embedding_function=get_embeddings(config["embedding_model"]),
            collection_name=config["collection_name"],
        )
        vectorstore.add_documents(chunks)

    updated_tracker = {
        filename: {
            "sha256": current_digests[filename],
            "chunk_count": (
                chunk_counts.get(filename, 0)
                if filename in newly_indexed
                else tracker[filename].get("chunk_count", 0)
            ),
        }
        for filename in current_files
    }
    save_tracker(updated_tracker)

    return vectorstore, newly_indexed, files_without_text
